"""
Document conversion module using Docling.

Converts unstructured documents (PDF, HTML, DOCX) to structured Markdown
with preserved hierarchy (headers, tables, code blocks).
"""

import json
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

logger = logging.getLogger(__name__)


@dataclass
class ConversionMetrics:
    """Metrics collected during document conversion."""
    docs_converted: int = 0
    docs_failed: int = 0
    chars_extracted: int = 0
    tables_detected: int = 0
    headers_found: int = 0
    code_blocks: int = 0
    conversion_time_seconds: float = 0.0
    errors: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "docs_converted": self.docs_converted,
            "docs_failed": self.docs_failed,
            "chars_extracted": self.chars_extracted,
            "tables_detected": self.tables_detected,
            "headers_found": self.headers_found,
            "code_blocks": self.code_blocks,
            "conversion_time_seconds": round(self.conversion_time_seconds, 2),
            "errors": self.errors,
        }


@dataclass
class ConvertedDocument:
    """A document converted to markdown with metadata."""
    source_path: Path
    product: str
    markdown: str
    title: str
    headers: list[dict]  # [{level: int, text: str, line: int}, ...]
    tables_count: int
    code_blocks_count: int
    char_count: int

    def to_dict(self) -> dict:
        return {
            "source_path": str(self.source_path),
            "product": self.product,
            "title": self.title,
            "headers": self.headers,
            "tables_count": self.tables_count,
            "code_blocks_count": self.code_blocks_count,
            "char_count": self.char_count,
        }


class DocumentConverter:
    """
    Converts PDF, HTML, and DOCX files to structured Markdown using Docling.

    Usage:
        converter = DocumentConverter()
        for doc in converter.convert_directory(Path("docs/product_docs/raw/pdf")):
            print(doc.title, doc.char_count)
    """

    def __init__(self, use_ocr: bool = True, ocr_lang: str = "en"):
        """
        Initialize the document converter.

        Args:
            use_ocr: Whether to use OCR for scanned documents
            ocr_lang: Language for OCR (default: English)
        """
        self.use_ocr = use_ocr
        self.ocr_lang = ocr_lang
        self._converter = None
        self.metrics = ConversionMetrics()

    def _get_converter(self):
        """Lazy initialization of Docling converter."""
        if self._converter is None:
            try:
                from docling.document_converter import DocumentConverter as DoclingConverter
                from docling.datamodel.pipeline_options import PdfPipelineOptions
                from docling.datamodel.base_models import InputFormat
                from docling.document_converter import PdfFormatOption

                pipeline_options = PdfPipelineOptions()
                pipeline_options.do_ocr = self.use_ocr

                self._converter = DoclingConverter(
                    format_options={
                        InputFormat.PDF: PdfFormatOption(pipeline_options=pipeline_options)
                    }
                )
            except ImportError as e:
                logger.warning(f"Docling not available: {e}. Using fallback converter.")
                self._converter = "fallback"
        return self._converter

    def convert_file(self, file_path: Path, product: str = "unknown") -> ConvertedDocument | None:
        """
        Convert a single file to markdown.

        Args:
            file_path: Path to the file (PDF, HTML, or DOCX)
            product: Product name for metadata

        Returns:
            ConvertedDocument or None if conversion fails
        """
        import time
        start_time = time.time()

        suffix = file_path.suffix.lower()

        try:
            if suffix == ".pdf":
                result = self._convert_pdf(file_path)
            elif suffix in (".html", ".htm"):
                result = self._convert_html(file_path)
            elif suffix == ".docx":
                result = self._convert_docx(file_path)
            elif suffix == ".md":
                # Already markdown, just read it
                result = self._read_markdown(file_path)
            else:
                logger.warning(f"Unsupported file type: {suffix} for {file_path}")
                return None

            if result is None:
                self.metrics.docs_failed += 1
                return None

            markdown, title = result

            # Extract structure from markdown
            headers = self._extract_headers(markdown)
            tables_count = markdown.count("|---") + markdown.count("| ---")
            code_blocks_count = markdown.count("```")

            elapsed = time.time() - start_time
            self.metrics.docs_converted += 1
            self.metrics.chars_extracted += len(markdown)
            self.metrics.headers_found += len(headers)
            self.metrics.tables_detected += tables_count
            self.metrics.code_blocks += code_blocks_count
            self.metrics.conversion_time_seconds += elapsed

            return ConvertedDocument(
                source_path=file_path,
                product=product,
                markdown=markdown,
                title=title or file_path.stem,
                headers=headers,
                tables_count=tables_count,
                code_blocks_count=code_blocks_count,
                char_count=len(markdown),
            )

        except Exception as e:
            logger.error(f"Failed to convert {file_path}: {e}")
            self.metrics.docs_failed += 1
            self.metrics.errors.append({"file": str(file_path), "error": str(e)})
            return None

    def _convert_pdf(self, file_path: Path) -> tuple[str, str] | None:
        """Convert PDF using Docling or fallback."""
        converter = self._get_converter()

        if converter == "fallback":
            return self._convert_pdf_fallback(file_path)

        try:
            result = converter.convert(str(file_path))
            markdown = result.document.export_to_markdown()
            title = getattr(result.document, 'title', None) or file_path.stem
            return markdown, title
        except Exception as e:
            logger.warning(f"Docling failed for {file_path}: {e}. Trying fallback.")
            return self._convert_pdf_fallback(file_path)

    def _convert_pdf_fallback(self, file_path: Path) -> tuple[str, str] | None:
        """Fallback PDF conversion using PyMuPDF."""
        try:
            import fitz  # PyMuPDF

            doc = fitz.open(file_path)
            markdown_parts = []
            title = doc.metadata.get("title", "") or file_path.stem

            for page_num, page in enumerate(doc):
                text = page.get_text("text")
                if text.strip():
                    markdown_parts.append(f"<!-- Page {page_num + 1} -->\n\n{text}")

            doc.close()
            return "\n\n".join(markdown_parts), title

        except ImportError:
            logger.error("PyMuPDF (fitz) not installed. Cannot convert PDF.")
            return None
        except Exception as e:
            logger.error(f"Fallback PDF conversion failed for {file_path}: {e}")
            return None

    def _convert_html(self, file_path: Path) -> tuple[str, str] | None:
        """Convert HTML to markdown."""
        converter = self._get_converter()

        if converter != "fallback":
            try:
                result = converter.convert(str(file_path))
                markdown = result.document.export_to_markdown()
                title = getattr(result.document, 'title', None) or file_path.stem
                return markdown, title
            except Exception as e:
                logger.warning(f"Docling failed for HTML {file_path}: {e}. Trying fallback.")

        # Fallback: use html2text or basic extraction
        return self._convert_html_fallback(file_path)

    def _convert_html_fallback(self, file_path: Path) -> tuple[str, str] | None:
        """Fallback HTML conversion using html2text or BeautifulSoup."""
        try:
            import html2text

            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                html_content = f.read()

            h = html2text.HTML2Text()
            h.ignore_links = False
            h.ignore_images = True
            h.body_width = 0  # No wrapping

            markdown = h.handle(html_content)

            # Try to extract title
            title = file_path.stem
            import re
            title_match = re.search(r'<title[^>]*>([^<]+)</title>', html_content, re.IGNORECASE)
            if title_match:
                title = title_match.group(1).strip()

            return markdown, title

        except ImportError:
            logger.warning("html2text not installed. Trying BeautifulSoup fallback.")
            return self._convert_html_beautifulsoup(file_path)
        except Exception as e:
            logger.error(f"html2text conversion failed for {file_path}: {e}")
            return self._convert_html_beautifulsoup(file_path)

    def _convert_html_beautifulsoup(self, file_path: Path) -> tuple[str, str] | None:
        """Last resort HTML conversion using BeautifulSoup."""
        try:
            from bs4 import BeautifulSoup

            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                soup = BeautifulSoup(f.read(), 'html.parser')

            # Remove script and style elements
            for element in soup(['script', 'style', 'nav', 'footer', 'header']):
                element.decompose()

            title = soup.title.string if soup.title else file_path.stem
            text = soup.get_text(separator='\n\n')

            # Clean up excessive whitespace
            import re
            text = re.sub(r'\n{3,}', '\n\n', text)
            text = re.sub(r' {2,}', ' ', text)

            return text.strip(), title

        except ImportError:
            logger.error("BeautifulSoup not installed. Cannot convert HTML.")
            return None
        except Exception as e:
            logger.error(f"BeautifulSoup conversion failed for {file_path}: {e}")
            return None

    def _convert_docx(self, file_path: Path) -> tuple[str, str] | None:
        """Convert DOCX to markdown."""
        converter = self._get_converter()

        if converter != "fallback":
            try:
                result = converter.convert(str(file_path))
                markdown = result.document.export_to_markdown()
                title = getattr(result.document, 'title', None) or file_path.stem
                return markdown, title
            except Exception as e:
                logger.warning(f"Docling failed for DOCX {file_path}: {e}. Trying fallback.")

        # Fallback using python-docx
        try:
            from docx import Document

            doc = Document(file_path)
            paragraphs = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    # Check for heading styles
                    if para.style and para.style.name.startswith('Heading'):
                        level = int(para.style.name[-1]) if para.style.name[-1].isdigit() else 1
                        paragraphs.append(f"{'#' * level} {text}")
                    else:
                        paragraphs.append(text)

            title = doc.core_properties.title or file_path.stem
            return "\n\n".join(paragraphs), title

        except ImportError:
            logger.error("python-docx not installed. Cannot convert DOCX.")
            return None
        except Exception as e:
            logger.error(f"DOCX conversion failed for {file_path}: {e}")
            return None

    def _read_markdown(self, file_path: Path) -> tuple[str, str] | None:
        """Read existing markdown file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()

            # Extract title from first H1 or filename
            title = file_path.stem
            for line in content.split('\n'):
                if line.startswith('# '):
                    title = line[2:].strip()
                    break

            return content, title
        except Exception as e:
            logger.error(f"Failed to read markdown {file_path}: {e}")
            return None

    def _extract_headers(self, markdown: str) -> list[dict]:
        """Extract header structure from markdown."""
        headers = []
        for line_num, line in enumerate(markdown.split('\n'), 1):
            line = line.strip()
            if line.startswith('#'):
                # Count the level
                level = 0
                for char in line:
                    if char == '#':
                        level += 1
                    else:
                        break
                text = line[level:].strip()
                if text:
                    headers.append({
                        "level": level,
                        "text": text,
                        "line": line_num,
                    })
        return headers

    def convert_directory(
        self,
        input_dir: Path,
        extensions: list[str] | None = None
    ) -> Iterator[ConvertedDocument]:
        """
        Convert all supported files in a directory.

        Args:
            input_dir: Directory to process
            extensions: File extensions to process (default: [".pdf", ".html", ".htm", ".docx", ".md"])

        Yields:
            ConvertedDocument for each successfully converted file
        """
        if extensions is None:
            extensions = [".pdf", ".html", ".htm", ".docx", ".md"]

        input_dir = Path(input_dir)

        # Find all files recursively
        for ext in extensions:
            for file_path in input_dir.rglob(f"*{ext}"):
                # Determine product from directory structure
                # e.g., docs/product_docs/raw/pdf/turbonomic/file.pdf -> turbonomic
                relative = file_path.relative_to(input_dir)
                product = relative.parts[0] if len(relative.parts) > 1 else "unknown"

                logger.info(f"Converting: {file_path}")
                doc = self.convert_file(file_path, product=product)
                if doc:
                    yield doc

    def save_markdown(self, doc: ConvertedDocument, output_dir: Path) -> Path:
        """
        Save converted document to output directory.

        Args:
            doc: Converted document
            output_dir: Output directory

        Returns:
            Path to saved file
        """
        output_dir = Path(output_dir)

        # Create product subdirectory
        product_dir = output_dir / doc.product
        product_dir.mkdir(parents=True, exist_ok=True)

        # Save with .md extension
        output_file = product_dir / f"{doc.source_path.stem}.md"

        # Add metadata header
        header = f"""---
source: {doc.source_path}
product: {doc.product}
title: {doc.title}
---

"""
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(header + doc.markdown)

        return output_file


def convert_documents(
    input_dir: Path,
    output_dir: Path,
    extensions: list[str] | None = None,
    use_ocr: bool = True,
) -> ConversionMetrics:
    """
    Convenience function to convert all documents in a directory.

    Args:
        input_dir: Input directory with raw documents
        output_dir: Output directory for markdown files
        extensions: File extensions to process
        use_ocr: Whether to use OCR for scanned PDFs

    Returns:
        ConversionMetrics with summary statistics
    """
    converter = DocumentConverter(use_ocr=use_ocr)
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    for doc in converter.convert_directory(input_dir, extensions):
        output_path = converter.save_markdown(doc, output_dir)
        logger.info(f"Saved: {output_path} ({doc.char_count} chars)")

    return converter.metrics


if __name__ == "__main__":
    import argparse

    logging.basicConfig(level=logging.INFO, format='%(levelname)s: %(message)s')

    parser = argparse.ArgumentParser(description="Convert documents to markdown")
    parser.add_argument("--input", "-i", required=True, help="Input directory")
    parser.add_argument("--output", "-o", required=True, help="Output directory")
    parser.add_argument("--no-ocr", action="store_true", help="Disable OCR")
    parser.add_argument("--extensions", "-e", nargs="+", help="File extensions to process")

    args = parser.parse_args()

    metrics = convert_documents(
        input_dir=Path(args.input),
        output_dir=Path(args.output),
        extensions=args.extensions,
        use_ocr=not args.no_ocr,
    )

    print("\nConversion Summary:")
    print(json.dumps(metrics.to_dict(), indent=2))
